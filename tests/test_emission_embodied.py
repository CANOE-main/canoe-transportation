"""Saved GREET extraction, road aggregation and shared insertion contracts."""

from dataclasses import replace
from pathlib import Path
import shutil
import json
import sqlite3

from canoe_schema.v4_0 import Region, TimePeriod
import openpyxl
import pandas as pd
from pydantic import ValidationError
import pytest
import yaml

from build_transport import insert_transport_contribution, prepare_transport_contribution
from fetching.greet_vehicle_cycle import (
    COMPONENT, SOURCE, GreetRegistration, extract_summary_rows, generation_profile,
    normalize_generated_results, registered_workbooks, required_inputs, result_bank_paths,
)
from fetching.greet_automation import inspect_lightweight_glider_formulas, probe_archetypes, probe_cases, publish_generated_results, validate_probe_results
from parameterization.road_embodied_emissions import prepare_emission_embodied_rows
from parameterization.road_fleet_weights import load_fleet_aggregation_evidence
from utils import file_sha256, load_config_bundle, load_harmonization_rules, resolve_artifact_path
from validation.config_models import ScenarioConfig
from validation.database_bootstrap import validate_database
from validation.insertion import insert_models, validate_transport_parameter_support
from validation.schema_contract import create_v4_schema


ROOT = Path(__file__).resolve().parents[1]
SCENARIO = "config/scenarios/legacy_reproduction.yaml"
LDV_CLASSES = {"Car": 1, "SUV": 2, "Pickup": 3}
TRUCK_CLASSES = {"Class 6 PnD Trucks": 1, "Class 8 Day-Cab Trucks": 2, "Class 8 Sleeper-Cab Trucks": 3}
LDV_BLOCKS = [
    "1.1) ICEV: Conventional Material", "1.2) ICEV: Lightweight Material",
    "1.3) HEV: Conventional Material", "1.4) HEV: Lightweight Material",
    "1.5) PHEV: Conventional Material", "1.6) PHEV: Lightweight Material",
    "1.7) EV: Conventional Material", "1.8) EV: Lightweight Material",
    "1.9) FCV: Conventional Material", "1.10) FCV: Lightweight Material",
]


def make_workbooks(directory: Path, ldv: str, truck: str) -> None:
    """Independent fixture layout with distinct combined CO2 and material values."""
    directory.mkdir(parents=True, exist_ok=True)
    first = openpyxl.Workbook()
    sheet = first.active
    sheet.title = "Inputs"
    sheet["B8"], sheet["E9"], sheet["E16"] = "1.1) Target Year for Simulation", 2025, LDV_CLASSES[ldv]
    first.save(directory / "R&D GREET1_2025_Rev1.xlsm")
    second = openpyxl.Workbook()
    second.remove(second.active)
    for name, code, header, labels in (
        ("Vehi_Inputs", LDV_CLASSES[ldv], "1. Selection of Vehicle Types for Simulation",
         ["1 -- Passenger Cars", "2 -- Sport Utility Vehicles", "3 -- Pick-Up Trucks"]),
        ("MHDV_Inputs", TRUCK_CLASSES[truck], "1. Selection of Truck Types for Simulation",
         ["1 -- Class 6 PnD Trucks", "2 -- Class 8 Day-Cab Trucks", "3 -- Class 8 Sleeper-Cab Trucks"]),
    ):
        sheet = second.create_sheet(name)
        sheet["A2"], sheet["B3"] = header, code
        for row, label in enumerate(labels, 3):
            sheet.cell(row, 3, label)
    sheet = second.create_sheet("GREET1_Import_Export")
    sheet["B4"], sheet["A20"] = LDV_CLASSES[ldv], "Vehicle Model Year"
    for cell in ("B20", "C20", "D20"):
        sheet[cell] = 2020
    for name, mass in (("Car", 2383.14), ("SUV", 2940.237), ("PUT", 3438.48112)):
        sheet = second.create_sheet(name)
        sheet["J4"] = "=IF(GREET1_Import_Export!$G$12=200,W4,IF(GREET1_Import_Export!$G$12=300,Y4,IF(GREET1_Import_Export!$G$12=400,AA4,AC4)))+(J17+IF(Vehi_Inputs!$C$76=1,J18,J19))"
        for column, miles in (("W", 150), ("Y", 200), ("AA", 300), ("AC", 400)):
            sheet[f"{column}3"] = f"EV{miles}: Lightweight Material"
            sheet[f"{column}4"] = mass
    for name, column, labels, code in (
        ("Vehi_Sum", "F", LDV_BLOCKS, LDV_CLASSES[ldv]),
        ("MHDV_Sum", "G", ["1.1) ICEV", "1.2) HEV", "1.3) EV", "1.4) FCV"], TRUCK_CLASSES[truck]),
    ):
        sheet = second.create_sheet(name)
        sheet["A1"] = "Energy Use and Emissions of Vehicle Cycle" + (": Trucks" if name == "MHDV_Sum" else "")
        for block, label in enumerate(labels):
            row = 13 + block * 31
            sheet[f"A{row-9}"] = label
            sheet[f"B{row-8}"] = "mmBtu or grams per vehicle lifetime"
            sheet[f"{column}{row-7}"] = "Total (No Credits)"
            multiplier = code * ((block // 2 + 1) if name == "Vehi_Sum" else (block + 1))
            if name == "Vehi_Sum" and block % 2:
                multiplier *= 2
            for offset, (gas, value) in enumerate((("CO2", 1e6), ("CO2 (VOC, CO, CO2)", 1e10), ("CH4", 1000), ("N2O", 10))):
                sheet[f"A{row+offset}"] = gas
                sheet[f"{column}{row+offset}"] = multiplier * value
    second.save(directory / "R&D GREET2_2025_Rev1.xlsm")


@pytest.fixture
def bundle(tmp_path):
    shutil.copytree(ROOT / "config", tmp_path / "config")
    shutil.copytree(ROOT / "inputs/0_canoe_template", tmp_path / "inputs/0_canoe_template")
    original = load_config_bundle(SCENARIO, repo_root=ROOT)
    fleet = load_fleet_aggregation_evidence(original)
    for path in fleet.paths:
        target = tmp_path / path.relative_to(ROOT)
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(path, target)
    scenario = tmp_path / SCENARIO
    payload = yaml.safe_load(scenario.read_text(encoding="utf-8"))
    payload["embodied_emissions"] = True
    payload["geography"]["regions"] = ["ON", "QC", "BCT", "NL", "PE"]
    scenario.write_text(yaml.safe_dump(payload, sort_keys=False), encoding="utf-8")
    # Preserve the explicit exclusion test; production includes approved CNG=ICEV.
    rules_path = tmp_path / "config/parameters/rules.yaml"
    payload = yaml.safe_load(rules_path.read_text(encoding="utf-8"))
    payload["parameterization"]["road_embodied_emissions"]["excluded_powertrains"] = ["cng"]
    rules_path.write_text(yaml.safe_dump(payload, sort_keys=False), encoding="utf-8")
    config = load_config_bundle(SCENARIO, repo_root=tmp_path)
    atb_rules = load_harmonization_rules(config, "nlr_atb_autonomie")
    atb_relative = Path(config.paths.inputs.interim) / atb_rules["interim_subdir"] / atb_rules["components"]["vehicles"]["output_file"]
    (tmp_path / atb_relative).parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(ROOT / atb_relative, tmp_path / atb_relative)
    registration, _ = registered_workbooks(config)
    make_workbooks(tmp_path / config.paths.inputs.external / registration.directory, "Car", "Class 6 PnD Trucks")
    write_fixture_bank(config)
    return config


def summary_cells(family, code, phev=0, bev=0):
    # Independent worksheet fixtures retain literal native anchors and labels.
    labels = LDV_BLOCKS if family == "ldv" else ["1.1) ICEV", "1.2) HEV", "1.3) EV", "1.4) FCV"]
    column = "F" if family == "ldv" else "G"
    cells = {"A1": "Energy Use and Emissions of Vehicle Cycle" + (": Trucks" if family == "mhdv" else "")}
    for block, label in enumerate(labels):
        row = 13 + block * 31
        cells[f"A{row-9}"] = label
        cells[f"B{row-8}"] = "mmBtu or grams per vehicle lifetime"
        cells[f"{column}{row-7}"] = "Total (No Credits)"
        multiplier = code * ((block // 2 + 1) if family == "ldv" else block + 1)
        distance = (phev if block in (4, 5) else bev if block in (6, 7) else 0) if family == "ldv" else 0
        multiplier += distance / 1000
        if family == "ldv" and block % 2:
            multiplier *= 2
        for offset, (gas, value) in enumerate((("CO2", 1e6), ("CO2 (VOC, CO, CO2)", 1e10), ("CH4", 1000), ("N2O", 10))):
            cells[f"A{row+offset}"] = gas
            cells[f"{column}{row+offset}"] = multiplier * value
    return cells


def write_fixture_bank(config):
    registration, books = registered_workbooks(config)
    settings = registration.automation
    rules = load_harmonization_rules(config, "greet_vehicle_cycle")
    ranges, atb = probe_archetypes(config)
    cases = probe_cases(rules, ranges)
    raw, excluded = [], []
    for case in cases:
        model = {**case, "solved_simulation_year": registration.solved_simulation_year,
                 "ldv_vehicle_cohort_year": registration.ldv_vehicle_cohort_year}
        for number, path in enumerate(books, 1):
            model[f"greet{number}_file"] = path.relative_to(config.repo_root).as_posix()
            model[f"greet{number}_sha256"] = file_sha256(path)
        for family in ("ldv", "mhdv"):
            cells = summary_cells(family, case[family + "_selector"], case["phev_range_miles"], case["bev_range_miles"])
            rows, exclusions = extract_summary_rows(cells, family=family, model=model, rules=rules)
            raw.extend(rows)
            excluded.extend(exclusions)
    formula_audit = inspect_lightweight_glider_formulas(books[1], settings.lightweight_glider)
    report = {"ok": True, "source_id": SOURCE, "component": COMPONENT,
              "profile_sha256": generation_profile(config), "source_hashes_unchanged": True,
              "source_hashes": {p.relative_to(config.repo_root).as_posix(): file_sha256(p) for p in books},
              "source_formula_audit": formula_audit,
              "lightweight_bev_range_mapping_ok": all(a["ok"] for a in formula_audit),
              "validation": validate_probe_results(pd.DataFrame(raw), settings, cases=cases, rules=rules)}
    run = config.repo_root / "fixture_generation"
    run.mkdir(exist_ok=True)
    pd.DataFrame(raw).to_csv(run / settings.files["results"], index=False)
    pd.DataFrame(excluded).to_csv(run / settings.files["exclusions"], index=False)
    pd.DataFrame([{**c, "converged_vehicle_cycle_totals": True, "calculation_passes": 2} for c in cases]).to_csv(run / settings.files["cases"], index=False)
    (run / settings.files["report"]).write_text(json.dumps(report))
    publish_generated_results(config, run)


def test_extraction_retains_exact_individual_gases_and_all_blocks(bundle):
    result = normalize_generated_results(bundle)
    assert len(result.rows) == 198
    assert len(result.exclusions) == 322
    assert set(result.rows.gas_label) == {"CO2", "CH4", "N2O"}
    assert set(result.exclusions.gas_label) == {"CO2 (VOC, CO, CO2)"}
    assert set(result.rows.loc[result.rows.family.eq("ldv"), "powertrain"]) == {"icev", "hev", "phev", "bev", "fcev"}
    assert set(result.rows.loc[result.rows.family.eq("ldv"), "materials"]) == {"conventional", "lightweight"}
    assert set(result.rows.loc[result.rows.family.eq("mhdv"), "materials"]) == {"none"}
    assert set(result.rows.solved_simulation_year) == {2025}
    assert set(result.rows.ldv_vehicle_cohort_year) == {2020}
    selected = result.rows.query("vehicle_class=='Car' and powertrain=='icev' and materials=='conventional' and gas_label=='CO2'").iloc[0]
    assert (selected.worksheet, selected.cell, selected.source_value) == ("Vehi_Sum", "F13", 1e6)
    last = result.rows.query("vehicle_class=='Car' and powertrain=='fcev' and materials=='lightweight' and gas_label=='N2O'").iloc[0]
    assert last.cell == "F295"


@pytest.mark.parametrize(("family", "cell", "value"), [
    ("ldv", "A13", "CO2 (VOC, CO, CO2)"), ("ldv", "A14", "CO2"),
    ("ldv", "A128", "Unexpected PHEV layout"), ("ldv", "B5", "g/mile"),
    ("ldv", "F6", "Total including credits"), ("ldv", "F13", None),
    ("ldv", "F13", "=1+1"), ("ldv", "F15", "#VALUE!"), ("mhdv", "G13", -1),
])
def test_extraction_rejects_unexpected_layout_units_or_missing_results(bundle, family, cell, value):
    cells = summary_cells(family, 1)
    cells[cell] = value
    with pytest.raises(ValueError, match="GREET"):
        extract_summary_rows(cells, family=family, model={"ldv_class": "Car", "mhdv_class": "Class 6 PnD Trucks"}, rules=load_harmonization_rules(bundle, "greet_vehicle_cycle"))


def test_source_sensitive_configuration_invalidates_existing_bank(bundle):
    bundle.sources.sources[SOURCE].component(COMPONENT).adapter["automation"]["names"]["bev_range_miles"]["cell"] = "Z7"
    with pytest.raises(ValueError, match="Stale GREET"):
        normalize_generated_results(bundle)


def test_registration_rejects_duplicate_classes_and_path_escape(bundle):
    payload = bundle.sources.sources[SOURCE].component(COMPONENT).adapter.copy()
    payload["greet2_file"] = "../workbook.xlsm"
    with pytest.raises(ValidationError, match="relative .xlsm"):
        GreetRegistration.model_validate(payload)
    payload = bundle.sources.sources[SOURCE].component(COMPONENT).adapter.copy()
    payload["greet2_file"] = payload["greet1_file"]
    with pytest.raises(ValidationError, match="Duplicate registered GREET"):
        GreetRegistration.model_validate(payload)


def test_disabled_requires_no_workbooks_or_preparation_io(bundle, monkeypatch):
    payload = bundle.scenario.model_dump()
    payload["embodied_emissions"] = False
    disabled = replace(bundle, scenario=ScenarioConfig.model_validate(payload))
    monkeypatch.setattr("parameterization.road_embodied_emissions.normalize_generated_results", lambda *_: pytest.fail("disabled extraction"))
    assert required_inputs(disabled) == []
    prepared = prepare_emission_embodied_rows(disabled)
    assert prepared.rows == prepared.provenance_contexts == []
    assert prepared.source_evidence is None
    assert prepared.audit == {"enabled": False, "rows": 0}
    assert not resolve_artifact_path(disabled, "emission_embodied_processed").exists()


def test_lifetime_conversion_renormalization_ownership_and_coverage(bundle):
    prepared = prepare_emission_embodied_rows(bundle)
    assert prepared.audit["ok"]
    assert prepared.audit["regions"] == ["BCT", "NLLAB", "ON", "PEI", "QC"]
    assert prepared.audit["vintages"] == [2025, 2030, 2035, 2040, 2045]
    assert len(prepared.rows) == 43 * 5 * 5 * 3
    assert all(r.units == ("kt/k vehicles" if r.emis_comm == "co2" else "t/k vehicles") for r in prepared.rows)
    assert prepared.audit["conversion_factors"] == {"CO2": 1e-6, "CH4": 1e-3, "N2O": 1e-3}
    assert prepared.audit["output_units"] == {"co2": "kt/k vehicles", "ch4": "t/k vehicles", "n2o": "t/k vehicles"}
    row = next(r for r in prepared.rows if r.region == "ON" and r.tech == "T_LDV_C_GSL_N" and r.emis_comm == "co2" and r.vintage == 2025)
    assert row.value == 1.0
    for gas, value in (("ch4", 1.0), ("n2o", .01)):
        gas_row = next(r for r in prepared.rows if r.region == "ON" and r.tech == "T_LDV_C_GSL_N" and r.emis_comm == gas and r.vintage == 2025)
        assert gas_row.value == value
    evidence = prepared.aggregation_evidence
    for gas, factor in (("CO2", 1e-6), ("CH4", 1e-3), ("N2O", 1e-3)):
        selected = evidence.loc[evidence.gas_label.eq(gas)]
        assert selected.conversion_factor.eq(factor).all()
        assert selected.weighted_value.tolist() == pytest.approx((selected.source_value * factor * selected.aggregation_weight).tolist())
    fleet = load_fleet_aggregation_evidence(bundle)
    weights = fleet.ldv.loc[fleet.ldv.nrcan_ceud_class.eq("Light Truck")].set_index("nlr_atb_class").aggregation_weight
    expected = (weights.Pickup * 3 + (weights["Small SUV"] + weights["Midsize SUV"]) * 2) / weights.sum()
    truck = next(r for r in prepared.rows if r.region == "ON" and r.tech == "T_LDV_LTP_GSL_N" and r.emis_comm == "co2")
    assert truck.value == pytest.approx(expected)
    assert all(r.emis_comm != "co2e" for r in prepared.rows)
    assert all(r.tech.endswith("_N") for r in prepared.rows)
    assert prepared.coverage.query("powertrain=='cng'").status.eq("unsupported_powertrain_explicitly_excluded").all()
    mhdv = prepared.aggregation_evidence.query("mode=='medium_trucks' and scenario_region=='ON' and tech=='T_MDV_T_DSL_N' and gas_label=='CO2'")
    assert set(mhdv.native_class) == {"2", "3", "4", "5", "6", "7"}
    assert mhdv.aggregation_weight.sum() == pytest.approx(1)
    assert set(mhdv.vehicle_class) == {"Class 6 PnD Trucks"}


def test_partial_pickup_suv_weights_are_renormalized_proportionately(bundle):
    fleet = load_fleet_aggregation_evidence(bundle)
    from parameterization.road_embodied_emissions import class_weights

    ldv = pd.concat([fleet.ldv, pd.DataFrame([{
        "nrcan_ceud_class": "Light Truck", "nlr_atb_class": "Unrepresented Van", "aggregation_weight": 0.5,
    }])], ignore_index=True)
    ldv.loc[ldv.nlr_atb_class.isin(["Pickup", "Small SUV", "Midsize SUV"]), "aggregation_weight"] *= 0.5
    altered = replace(fleet, ldv=ldv)
    rules = load_harmonization_rules(bundle, "road_embodied_emissions")
    adapted = class_weights(altered, "ON", "passenger_light_trucks", rules=rules)
    assert sum(w["original_weight"] for w in adapted) == pytest.approx(0.5)
    assert sum(w["aggregation_weight"] for w in adapted) == pytest.approx(1)
    assert all(w["aggregation_weight"] == pytest.approx(w["original_weight"] / 0.5) for w in adapted)


def test_both_materials_keep_mhdv_rows_and_provenance_identical(bundle):
    conventional = prepare_emission_embodied_rows(bundle)
    payload = bundle.scenario.model_dump()
    payload["embodied_materials"] = "lightweight"
    lightweight = prepare_emission_embodied_rows(replace(bundle, scenario=ScenarioConfig.model_validate(payload)))
    assert not lightweight.source_evidence.audit["lightweight_bev_range_mapping_ok"]
    assert all(a["glider_inputs_equal"] for a in lightweight.source_evidence.audit["lightweight_glider_input_audit"])
    def mhd(prepared):
        return [r.model_dump() for r in prepared.rows if r.tech.startswith(("T_MDV", "T_HDV"))]

    assert mhd(conventional) == mhd(lightweight)
    def ldv(prepared):
        return [r.value for r in prepared.rows if r.tech.startswith("T_LDV")]

    assert ldv(lightweight) == pytest.approx([2 * value for value in ldv(conventional)])


def test_offline_preparation_is_deterministic_and_read_only(bundle, monkeypatch):
    import socket

    monkeypatch.setattr(socket, "create_connection", lambda *_args, **_kwargs: pytest.fail("network access"))
    monkeypatch.setattr("fetching.greet_automation.generate_vehicle_cycle_results", lambda *_: pytest.fail("Excel startup"))
    paths = registered_workbooks(bundle)[1]
    before = {p: p.read_bytes() for p in paths}
    first = prepare_emission_embodied_rows(bundle)
    dirs = [resolve_artifact_path(bundle, key) for key in ("greet_vehicle_cycle_interim", "emission_embodied_processed", "emission_embodied_validation")]
    artifacts = {p: p.read_bytes() for d in dirs for p in d.iterdir()}
    second = prepare_emission_embodied_rows(bundle)
    assert [r.model_dump() for r in first.rows] == [r.model_dump() for r in second.rows]
    assert first.provenance_contexts == second.provenance_contexts
    assert first.audit == second.audit
    assert all(p.read_bytes() == data for p, data in artifacts.items())
    assert all(p.read_bytes() == data for p, data in before.items())


@pytest.mark.parametrize("problem", ["missing_report", "altered_csv", "source_change", "wrong_anchor", "combined_co2"])
def test_complete_result_bank_rejects_missing_stale_or_invalid_evidence(bundle, problem):
    bank = result_bank_paths(bundle)
    if problem == "missing_report":
        bank["report"].unlink()
        expected = FileNotFoundError
    elif problem == "source_change":
        registered_workbooks(bundle)[1][1].write_bytes(b"changed registered source")
        expected = ValueError
    else:
        frame = pd.read_csv(bank["results"], float_precision="round_trip")
        if problem == "wrong_anchor":
            frame.loc[0, "cell"] = "F14"
        elif problem == "combined_co2":
            frame.loc[0, "gas_label"] = "CO2 (VOC, CO, CO2)"
        else:
            frame.loc[0, "source_value"] += 1
        frame.to_csv(bank["results"], index=False)
        if problem != "altered_csv":
            manifest = json.loads(bank["manifest"].read_text())
            manifest["files"][bank["results"].name] = file_sha256(bank["results"])
            bank["manifest"].write_text(json.dumps(manifest))
        expected = ValueError
    with pytest.raises(expected, match="GREET"):
        normalize_generated_results(bundle)


@pytest.mark.parametrize("corrected_selector", [False, True])
def test_unequal_glider_inputs_require_correct_selector_for_lightweight(bundle, corrected_selector):
    path = registered_workbooks(bundle)[1][1]
    book = openpyxl.load_workbook(path)
    sheet = book["SUV"]
    sheet["Y4"] = 3000
    if corrected_selector:
        sheet["J4"] = sheet["J4"].value.replace("=200,W4", "=150,W4").replace("=300,Y4", "=200,Y4").replace("=400,AA4", "=300,AA4")
    book.save(path)
    book.close()
    write_fixture_bank(bundle)
    assert prepare_emission_embodied_rows(bundle).rows
    payload = bundle.scenario.model_dump()
    payload["embodied_materials"] = "lightweight"
    lightweight = replace(bundle, scenario=ScenarioConfig.model_validate(payload))
    if corrected_selector:
        assert prepare_emission_embodied_rows(lightweight).rows
    else:
        with pytest.raises(ValueError, match="lightweight BEV glider-range mapping has unequal inputs"):
            prepare_emission_embodied_rows(lightweight)


def test_each_atb_range_uses_its_own_total_and_vintages_hold_constant(bundle):
    prepared = prepare_emission_embodied_rows(bundle)
    cars = [r for r in prepared.rows if r.region == "ON" and r.tech.startswith("T_LDV_C_BEV") and r.emis_comm == "co2"]
    assert cars
    by_tech = pd.DataFrame([r.model_dump() for r in cars]).groupby("tech").value
    assert by_tech.nunique().eq(1).all()
    assert len(set(by_tech.first())) == 4
    assert set(prepared.aggregation_evidence.query("powertrain=='bev' and mode=='cars'").range_miles) == {150, 200, 300, 400}


def prepared_contribution(bundle, connection):
    return prepare_transport_contribution(
        connection, bundle=bundle, template_dir=bundle.repo_root / "inputs/0_canoe_template",
        include_existing_capacity=False, include_demand=False, include_road_utilization=False,
        include_lifetimes=False, include_efficiencies=False, include_costs=False, include_ev_chargers=False,
    )


def test_shared_contribution_preserves_full_preparation_and_caller_transaction(bundle):
    connection = sqlite3.connect(":memory:")
    create_v4_schema(connection)
    insert_models(connection, [Region(region=r) for r in ("ON", "QC", "BCT", "NLLAB", "PEI")])
    insert_models(connection, [TimePeriod(period=p, flag="f") for p in bundle.scenario.periods.model])
    connection.commit()
    contribution = prepared_contribution(bundle, connection)
    assert not connection.in_transaction
    assert connection.execute("SELECT count(*) FROM emission_embodied").fetchone()[0] == 0
    assert contribution.emission_embodied is not None
    assert len(contribution.emission_embodied.source_evidence.rows) == 198
    assert not contribution.emission_embodied.aggregation_evidence.empty
    inserted = insert_transport_contribution(connection, contribution)
    assert connection.in_transaction
    keys = {table: [tuple(getattr(r, f) for f in r.__primary_key__) for r in rows] for table, rows in inserted.items()}
    assert validate_database(connection, expected_primary_keys=keys)["ok"]
    assert connection.execute("SELECT count(*) FROM emission_embodied WHERE dq_time IS NULL OR data_source IS NULL").fetchone()[0] == 0
    ldv_data_id = next(r.data_id for r in contribution.emission_embodied_rows if r.tech.startswith("T_LDV_C"))
    assert "Vehi_Sum" in connection.execute(
        "SELECT description FROM data_set WHERE data_id = ?", (ldv_data_id,),
    ).fetchone()[0]
    insert_transport_contribution(connection, contribution, conflict="ignore_identical")
    connection.rollback()
    assert connection.execute("SELECT count(*) FROM emission_embodied").fetchone()[0] == 0
    assert connection.execute("SELECT count(*) FROM data_set").fetchone()[0] == 0
    connection.close()


@pytest.mark.parametrize("problem", ["value", "row_units", "commodity_units"])
def test_preinsert_rechecks_mutated_embodied_rows_before_writes(bundle, problem):
    connection = sqlite3.connect(":memory:")
    create_v4_schema(connection)
    contribution = prepared_contribution(bundle, connection)
    if problem == "value":
        contribution.emission_embodied_rows[0].value = float("inf")
    elif problem == "row_units":
        row = next(r for r in contribution.emission_embodied_rows if r.emis_comm == "ch4")
        row.units = "kt/k vehicles"
    else:
        next(r for r in contribution.rows_by_table["commodity"] if r.name == "n2o").units = "kt"
    with pytest.raises(ValueError, match="Invalid embodied"):
        insert_transport_contribution(connection, contribution)
    assert connection.execute("SELECT count(*) FROM data_set").fetchone()[0] == 0
    connection.close()


def test_embodied_keys_require_efficiency_when_that_layer_is_prepared():
    from types import SimpleNamespace

    row = SimpleNamespace(region="ON", tech="T_NEW", vintage=2025)
    report = validate_transport_parameter_support({"emission_embodied": [row], "efficiency": []}, existing_vintages=[2020])
    assert not report["ok"]
    assert report["checks"]["emission_embodied_efficiency"]["unsupported_keys"] == 1


@pytest.mark.parametrize("problem", ["conversion", "trace_gas_conversion", "gas_owner", "gas_units", "parameter_units", "powertrain", "medium_class"])
def test_invalid_harmonization_or_owner_contract_is_rejected(bundle, problem):
    if problem in {"conversion", "trace_gas_conversion"}:
        path = bundle.repo_root / "config/parameters/conversion.yaml"
        payload = yaml.safe_load(path.read_text(encoding="utf-8"))
        key = "grams_per_vehicle_to_kt_per_thousand_vehicles" if problem == "conversion" else "grams_per_vehicle_to_tonnes_per_thousand_vehicles"
        payload["mass"][key] = 1e-9
        path.write_text(yaml.safe_dump(payload, sort_keys=False), encoding="utf-8")
        expected = "lifetime mass conversion/units"
    elif problem in {"gas_owner", "gas_units"}:
        path = bundle.repo_root / "inputs/0_canoe_template/commodity.csv"
        frame = pd.read_csv(path)
        if problem == "gas_owner":
            frame.loc[frame.name.eq("co2"), "flag"] = "p"
        else:
            frame.loc[frame.name.eq("ch4"), "units"] = "kt"
        frame.to_csv(path, index=False)
        expected = "emission commodity"
    else:
        path = bundle.repo_root / "config/parameters/rules.yaml"
        payload = yaml.safe_load(path.read_text(encoding="utf-8"))
        rules = payload["parameterization"]["road_embodied_emissions"]
        if problem == "parameter_units":
            rules["gas_scaling"]["N2O"]["output_units"] = "kt/k vehicles"
            expected = "lifetime mass conversion/units"
        elif problem == "powertrain":
            rules["excluded_powertrains"] = []
            rules["powertrain_map"]["ldv"].pop("cng")
            expected = "Unsupported GREET embodied powertrain"
        else:
            rules["medium_class_map"].pop(2)
            expected = "Unreviewed GREET medium-truck classes"
        path.write_text(yaml.safe_dump(payload, sort_keys=False), encoding="utf-8")
    with pytest.raises(ValueError, match=expected):
        prepare_emission_embodied_rows(bundle)
    assert not resolve_artifact_path(bundle, "emission_embodied_processed").exists()


def test_approved_cng_uses_icev_and_existing_vehicles_receive_no_rows(bundle):
    path = bundle.repo_root / "config/parameters/rules.yaml"
    payload = yaml.safe_load(path.read_text(encoding="utf-8"))
    payload["parameterization"]["road_embodied_emissions"].pop("excluded_powertrains")
    path.write_text(yaml.safe_dump(payload, sort_keys=False), encoding="utf-8")
    prepared = prepare_emission_embodied_rows(bundle)
    assert all(row.tech.endswith("_N") for row in prepared.rows)
    lookup = {(r.region, r.tech, r.emis_comm, r.vintage): r.value for r in prepared.rows}
    cng = [row for row in prepared.rows if "_CNG_" in row.tech]
    assert cng
    for row in cng:
        assert row.value == lookup[row.region, row.tech.replace("_CNG_", "_GSL_"), row.emis_comm, row.vintage]


def test_invalid_material_selection_and_missing_controls_are_rejected(bundle):
    payload = bundle.scenario.model_dump()
    payload["embodied_materials"] = "medium_truck_lightweight"
    with pytest.raises(ValidationError, match="embodied_materials"):
        ScenarioConfig.model_validate(payload)
    payload.pop("embodied_materials")
    with pytest.raises(ValidationError, match="embodied_materials"):
        ScenarioConfig.model_validate(payload)


def test_empty_legacy_embodied_baseline_is_reported_without_accepting_new_rows(tmp_path):
    from validation.legacy_compare import compare_legacy_emission_embodied

    candidate, reference = tmp_path / "candidate.sqlite", tmp_path / "reference.sqlite"
    with sqlite3.connect(candidate) as connection:
        connection.execute("CREATE TABLE emission_embodied (region, tech, emis_comm, vintage, value, units)")
        connection.execute("INSERT INTO emission_embodied VALUES ('ON', 'T_NEW', 'co2', 2025, 5, 'kt/k vehicles')")
    with sqlite3.connect(reference) as connection:
        connection.execute("CREATE TABLE EmissionEmbodied (region, tech, emis_comm, vintage, value, units)")
    report = compare_legacy_emission_embodied(candidate, reference, absolute_tolerance=1e-9, relative_tolerance=1e-6)
    assert report["candidate_rows"] == report["candidate_only_keys"] == 1
    assert report["reference_rows"] == report["shared_keys"] == 0
    assert "diagnostic only" in report["status"]
