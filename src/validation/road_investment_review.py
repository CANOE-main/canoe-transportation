"""Opt-in CAPEX evidence owned by validation.transport_comparison.

Products belong to the configured transport_comparison_report family. This module
reads registered/current audit evidence and preserved legacy worksheets; it never
prepares production parameters, modifies assumptions, or writes a database.
"""

from dataclasses import dataclass
from functools import cache
import json
from math import isclose, isfinite
from typing import Self
import warnings
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pandas as pd
from pydantic import Field, model_validator

from fetching.nlr_atb_autonomie import configured_trajectory
from parameterization.road_capex_opex import aggregate_purchase_price, select_nlr_purchase_prices
from parameterization.road_efficiencies import interpolate
from parameterization.road_fleet_weights import medium_vocation_weights
from utils import (
    ConfigBundle, file_sha256, load_harmonization_rules, resolve_artifact_path,
    resolve_input_path, resolve_repo_path,
)
from validation.config_models import MappingModel


ATB = "nlr_atb_transportation"


class CostReviewFinding(MappingModel):
    status: str = Field(min_length=1)
    driver: str = Field(min_length=1)
    scope: str = Field(min_length=1)
    interpretation: str = Field(min_length=1)
    evidence: list[str]


class RoadInvestmentReviewConfig(MappingModel):
    region: str = Field(min_length=1)
    legacy_workbook: str
    legacy_sheet: str
    legacy_currency_sheet: str
    legacy_currency_cell: str
    legacy_ldv_raw_workbook: str
    legacy_ldv_raw_sheet: str
    legacy_ldv_filters: dict[str, str]
    legacy_ldv_columns: dict[str, str]
    legacy_ldv_powertrains: dict[str, str]
    ldv_alternative_powertrains: dict[str, str]
    legacy_weight_cells: dict[str, dict[str, str | float]]
    legacy_medium_gvwr_cells: dict[int, str]
    legacy_heavy_groups: dict[str, list[str]]
    atb_inputs_member: str
    primary_sources: dict[str, str]
    findings: list[CostReviewFinding]

    @model_validator(mode="after")
    def complete_basis(self) -> Self:
        if set(self.legacy_ldv_filters) != {"performance", "progress"}:
            raise ValueError("Legacy LDV review requires performance and progress selections")
        if not {"performance", "progress", "class", "powertrain", "year", "manufacturing"} <= set(self.legacy_ldv_columns):
            raise ValueError("Legacy LDV source column contract is incomplete")
        return self


@dataclass(frozen=True)
class RoadInvestmentReview:
    attribution: pd.DataFrame
    weights: pd.DataFrame
    ldv_architecture: pd.DataFrame
    hydrogen_inputs: pd.DataFrame
    findings: pd.DataFrame
    metadata: dict

    def frames(self) -> dict[str, pd.DataFrame]:
        return {
            "road_cost_attribution.csv": self.attribution,
            "road_cost_weights.csv": self.weights,
            "road_cost_ldv_architecture.csv": self.ldv_architecture,
            "road_cost_hydrogen_inputs.csv": self.hydrogen_inputs,
            "road_cost_findings.csv": self.findings,
        }


def factor_identity(
    *, legacy: float, current: float, currency: float,
    weights: float, native_basis: float,
) -> float:
    """Reject a claimed multiplicative attribution that fails to reconcile."""
    values = (legacy, current, currency, weights, native_basis)
    if any(not isfinite(value) or value <= 0 for value in values):
        raise ValueError("CAPEX factor attribution requires finite positive values")
    residual = legacy * currency * weights * native_basis - current
    if not isclose(residual, 0, abs_tol=1e-8):
        raise ValueError("Road CAPEX factor attribution does not reproduce the compiled cost")
    return residual


def _legacy_weights(sheet, config, mode, current_weights, fleet_rules):
    if mode == "medium_trucks":
        shares = {gvwr: float(sheet[cell]) for gvwr, cell in config.legacy_medium_gvwr_cells.items()}
        weights = medium_vocation_weights(shares, list(current_weights), rules=fleet_rules["medium_trucks"])
        basis = "Legacy Ontario 2022 GVWR shares; current available freight vocations (Class 3 Box absent)"
    elif mode == "heavy_trucks":
        weights = {label: float(sheet[cell]) / len(labels)
                   for cell, labels in config.legacy_heavy_groups.items() for label in labels}
        basis = "Legacy 2018 shipment shares; equal daycab vocations"
    else:
        weights = {label: float(sheet[value]) if isinstance(value, str) else value
                   for label, value in config.legacy_weight_cells[mode].items()}
        basis = "Actual legacy worksheet weights" if mode != "school_buses" else "Actual legacy school formula reuses heavy-truck shipment weights"
    if set(weights) != set(current_weights) or any(value < 0 for value in weights.values()) or not isclose(sum(weights.values()), 1, abs_tol=1e-8):
        raise ValueError(f"Unsupported legacy/current class alignment: {mode}")
    return weights, basis


def _read_legacy(config, root):
    path = resolve_repo_path(root, config.legacy_workbook)
    # This warning concerns saving unsupported validation widgets. The workbook
    # is opened read-only, never saved, and its unchanged hash is review evidence.
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported and will be removed", category=UserWarning)
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = book[config.legacy_sheet]
            sheet = {f"{get_column_letter(col)}{index}": value
                     for index, row in enumerate(worksheet.iter_rows(max_col=44, values_only=True), 1)
                     for col, value in enumerate(row, 1) if value is not None}
            if sheet.get("A2") != "Technology" or sheet.get("B2") != "Region":
                raise ValueError("Unsupported legacy investment worksheet layout")
            technologies = {sheet[f"A{row}"]: row for row in range(3, worksheet.max_row + 1)
                            if str(sheet.get(f"A{row}", "")).startswith("T_")}
            years = {int(sheet[f"{get_column_letter(col)}2"]): get_column_letter(col)
                     for col in range(11, 18)}
            factor = float(book[config.legacy_currency_sheet][config.legacy_currency_cell].value)
        finally:
            book.close()
    return sheet, technologies, years, factor


def _raw_ldv(config, root):
    columns = config.legacy_ldv_columns
    frame = pd.read_excel(
        resolve_repo_path(root, config.legacy_ldv_raw_workbook),
        sheet_name=config.legacy_ldv_raw_sheet, usecols=list(columns.values()),
    ).rename(columns={value: key for key, value in columns.items()})
    frame["source_excel_row"] = frame.index + 2
    for field, value in config.legacy_ldv_filters.items():
        frame = frame.loc[frame[field].eq(value)]
    metrics = {name: "mean" for name in columns if name not in {
        "performance", "progress", "class", "powertrain", "year",
    }}
    metrics["source_excel_row"] = lambda values: ";".join(str(int(value)) for value in values)
    return frame.groupby(["class", "powertrain", "year"], as_index=False).agg(metrics)


def _source_checks(audit, compiled, retail, applied_rpe):
    if not isclose(float(audit.rpe_markup), applied_rpe, abs_tol=1e-12):
        raise ValueError("Cost audit uses a stale ATB family markup; rebuild before attribution")
    if not isclose(float(audit.cost), float(compiled.candidate), abs_tol=1e-8, rel_tol=1e-10):
        raise ValueError(f"Cost audit is stale for {audit.tech}/{audit.vintage}; rebuild before attribution")
    if not isclose(retail / applied_rpe, float(audit.source_value), abs_tol=1e-6, rel_tol=1e-10):
        raise ValueError(f"Selected ATB prices/weights do not reproduce the cost audit: {audit.tech}/{audit.vintage}")
    factor = float(audit.cad_per_usd) * float(audit.target_year_deflator) / float(audit.source_year_deflator)
    if not isclose(float(audit.source_value) * factor / float(audit.normalization_divisor),
                   float(audit.cost), abs_tol=1e-8, rel_tol=1e-10):
        raise ValueError("Currency audit does not reproduce compiled road CAPEX")
    if float(audit.normalization_divisor) != 1000:
        raise ValueError("Road CAPEX review requires per-vehicle to M$/thousand-vehicle units")
    return factor


def prepare_road_investment_review(
    bundle: ConfigBundle, comparison_rows: pd.DataFrame, config: RoadInvestmentReviewConfig,
) -> RoadInvestmentReview:
    """Refresh measured factors, controlled weight changes and source evidence."""
    costs = load_harmonization_rules(bundle, "costs")
    efficiency = load_harmonization_rules(bundle, "efficiencies")
    fleet_rules = load_harmonization_rules(bundle, "road_aggregation")
    atb_rules = load_harmonization_rules(bundle, "nlr_atb_autonomie")
    audit_path = resolve_artifact_path(bundle, "costs_interim", costs["files"]["road_invest"])
    price_path = resolve_input_path(bundle, "interim", atb_rules["interim_subdir"],
                                    atb_rules["components"]["vehicles"]["output_file"])
    technology_path = resolve_input_path(bundle, "template", "technology.csv")
    audit = pd.read_csv(audit_path, low_memory=False)
    audit = audit.loc[audit.region.eq(config.region) & audit.aggregation_weights.notna() & audit.rpe_applied.eq(True)]
    metadata = bundle.sources.sources[ATB].components["vehicles"].adapter
    prices = select_nlr_purchase_prices(
        pd.read_csv(price_path), trajectory=configured_trajectory(bundle, parameter="costs"),
        efficiency_rules=efficiency, rpe_markups=costs["nlr_road_retail_price_equivalent_markup"],
        source_currency=str(metadata["native_cost_currency"]),
        source_dollar_year=int(metadata["native_cost_dollar_year"]),
    )
    technology = pd.read_csv(technology_path).set_index("tech", verify_integrity=True)
    compiled = comparison_rows.loc[
        comparison_rows.table.eq("cost_invest") & comparison_rows.region.eq(config.region)
        & comparison_rows.status.isin(["equal", "different"])
    ].set_index(["tech", "vintage"], verify_integrity=True)
    sheet, legacy_tech, year_columns, legacy_factor = _read_legacy(config, bundle.repo_root)
    raw = _raw_ldv(config, bundle.repo_root)

    @cache
    def retail_price(powertrain, year, weights):
        return aggregate_purchase_price(prices, powertrain=powertrain, year=year, weights=dict(weights))[0]

    @cache
    def raw_value(vehicle_class, powertrain, year, metric="manufacturing"):
        selected = raw.loc[raw["class"].eq(vehicle_class) & raw.powertrain.eq(powertrain)]
        return interpolate(selected, year, metric)

    records, weight_records, architecture = [], {}, []
    for row in audit.itertuples(index=False):
        vintage, source_year = int(row.vintage), int(row.purchase_evidence_year)
        key = row.tech, vintage
        if key not in compiled.index:
            continue
        observed = compiled.loc[key]
        current_weights = json.loads(row.aggregation_weights)
        old_weights, weight_basis = _legacy_weights(sheet, config, row.mode, current_weights, fleet_rules)
        for label, value in current_weights.items():
            weight_records[row.mode, label] = {
                "mode": row.mode, "vehicle_class": label, "legacy_weight": old_weights[label],
                "current_weight": value, "change_percentage_points": 100 * (value - old_weights[label]),
                "legacy_basis": weight_basis,
            }
        family = "ldv" if efficiency["road_classes"][row.mode]["pathway"] == "ldv" else "mhdv"
        powertrain = efficiency["subcategory_aliases"].get(technology.at[row.tech, "sub_category"], technology.at[row.tech, "sub_category"])
        selected_powertrain = costs["missing_mhdv_purchase_powertrain_proxy"].get(powertrain, powertrain) if family == "mhdv" else powertrain
        current_retail = retail_price(selected_powertrain, source_year, tuple(sorted(current_weights.items())))
        old_weight_retail = retail_price(selected_powertrain, source_year, tuple(sorted(old_weights.items())))
        applied = float(costs["nlr_road_retail_price_equivalent_markup"][family])
        current_factor = _source_checks(row, observed, current_retail, applied)
        legacy_row = legacy_tech[row.tech]
        legacy_cell = f"{year_columns[vintage]}{legacy_row}"
        cached_cost = float(sheet[legacy_cell])
        if sheet[f"B{legacy_row}"] != config.region or not isclose(cached_cost, observed.legacy, abs_tol=1e-8, rel_tol=1e-10):
            raise ValueError(f"Legacy workbook does not reproduce the selected SQLite: {row.tech}/{vintage}")
        if not isclose(float(sheet[f"G{legacy_row}"]), legacy_factor, abs_tol=1e-12):
            raise ValueError("Legacy ATB road rows use inconsistent monetary conversion factors")
        legacy_native = cached_cost * 1000 / legacy_factor
        currency_factor = current_factor / legacy_factor
        weight_factor = current_retail / old_weight_retail
        native_factor = (old_weight_retail / applied) / legacy_native
        residual = factor_identity(
            legacy=float(observed.legacy), current=float(row.cost),
            currency=currency_factor, weights=weight_factor, native_basis=native_factor,
        )
        record = {
            "region": row.region, "mode": row.mode, "family": family, "tech": row.tech,
            "powertrain": powertrain, "source_powertrain": selected_powertrain,
            "vintage": vintage, "purchase_evidence_year": source_year,
            "legacy_cost": float(observed.legacy), "current_cost": float(row.cost),
            "gap_percent": 100 * (row.cost / observed.legacy - 1),
            "applied_rpe": applied, "currency_factor": currency_factor,
            "weight_factor": weight_factor, "native_basis_factor": native_factor,
            "current_conversion": current_factor, "legacy_conversion": legacy_factor,
            "current_dollar_year": int(row.source_dollar_year),
            "legacy_declared_dollar_year": int(sheet[f"E{legacy_row}"]),
            "current_native_manufacturing": current_retail / applied,
            "legacy_native_manufacturing": legacy_native,
            "legacy_weight_effect_percent": 100 * (old_weight_retail / current_retail - 1),
            "cost_with_legacy_weights": row.cost * old_weight_retail / current_retail,
            "reconciliation_residual": residual,
            "legacy_workbook_residual": cached_cost - observed.legacy,
            "legacy_output_cell": f"{config.legacy_sheet}!{legacy_cell}",
            "legacy_weight_basis": weight_basis, "units": row.units,
            "raw_ldv_reconstruction_residual": None,
        }
        if family == "ldv":
            original = config.legacy_ldv_powertrains[powertrain]
            alternative = config.ldv_alternative_powertrains.get(powertrain, original)
            legacy_raw_total = sum(old_weights[label] * raw_value(label, original, source_year) for label in old_weights)
            reconstructed = legacy_raw_total * legacy_factor / 1000
            if not isclose(reconstructed, cached_cost, abs_tol=1e-7, rel_tol=1e-10):
                raise ValueError(f"Raw legacy LDV prices do not reproduce the workbook: {row.tech}/{vintage}")
            record["raw_ldv_reconstruction_residual"] = reconstructed - observed.legacy
            for label in old_weights:
                original_native = raw_value(label, original, source_year)
                alternative_native = raw_value(label, alternative, source_year)
                selected_native = retail_price(selected_powertrain, source_year, ((label, 1.0),)) / applied
                architecture.append({
                    "mode": row.mode, "tech": row.tech, "vehicle_class": label,
                    "vintage": vintage, "legacy_powertrain": original,
                    "same_source_alternative": alternative,
                    "legacy_raw_manufacturing": original_native,
                    "alternative_raw_manufacturing": alternative_native,
                    "atb_manufacturing": selected_native,
                    "architecture_factor": alternative_native / original_native,
                    "atb_vs_same_architecture_factor": selected_native / alternative_native,
                    "native_total_factor": selected_native / original_native,
                    "legacy_weight": old_weights[label], "current_weight": current_weights[label],
                })
        records.append(record)
    attribution = pd.DataFrame(records)
    if attribution.empty:
        raise ValueError("No matched ATB-backed road investment rows available for review")
    attribution["gap_with_legacy_weights_percent"] = 100 * (
        attribution.cost_with_legacy_weights / attribution.legacy_cost - 1
    )
    source = bundle.sources.sources[ATB]
    archive_path = resolve_input_path(bundle, "cache", str(source.adapter["cache_path"]))
    with ZipFile(archive_path) as archive:
        inputs = pd.read_csv(archive.open(config.atb_inputs_member), low_memory=False)
    inputs = inputs.loc[
        inputs.scenario.eq(configured_trajectory(bundle, parameter="costs"))
        & inputs.vehicle_weight_category.eq(efficiency["atb"]["weight_categories"]["ldv"])
        & inputs.vehicle_detail.eq(efficiency["ldv_archetypes"]["fcev"]["detail"])
        & inputs.year.ge(min(bundle.scenario.periods.model))
    ]
    hydrogen = []
    for row in inputs.itertuples(index=False, name=None):
        item = dict(zip(inputs.columns, row, strict=True))
        label, year = item["vehicle_class"], int(item["year"])
        selected_output = retail_price("fcev", year, ((label, 1.0),))
        if not isclose(selected_output, float(item["vehicle_cost($)"]), abs_tol=1e-6, rel_tol=1e-10):
            raise ValueError(f"ATB input/output fuel-cell prices disagree: {label}/{year}")
        old_power = raw_value(label, config.legacy_ldv_powertrains["fcev"], year, "fuel_cell_power_w") / 1000
        old_stack = raw_value(label, config.legacy_ldv_powertrains["fcev"], year, "fuel_cell_cost")
        hydrogen.append({
            "vehicle_class": label, "source_year": year,
            "legacy_stack_cost_per_kw": old_stack / old_power,
            "atb_stack_cost_per_kw": item["fuel_cell_cost($/kW)"],
            "legacy_hydrogen_storage_total": raw_value(label, config.legacy_ldv_powertrains["fcev"], year, "hydrogen_storage_cost"),
            "legacy_usable_hydrogen_kg": raw_value(label, config.legacy_ldv_powertrains["fcev"], year, "usable_hydrogen_kg"),
            "atb_storage_cost_per_kg": item["h2_storage_tank_cost($/kgh2)"],
            "atb_declared_dollar_year": item["dollar_year"],
            "legacy_component_basis": "Raw fuel-cell header says 2020$, description/vehicle total says 2023$; storage says 2023$",
        })
    files = [audit_path, price_path, technology_path, archive_path,
             resolve_repo_path(bundle.repo_root, config.legacy_workbook),
             resolve_repo_path(bundle.repo_root, config.legacy_ldv_raw_workbook)]
    findings = pd.DataFrame([{**finding.model_dump(), "evidence": "; ".join(finding.evidence)} for finding in config.findings])
    return RoadInvestmentReview(
        attribution.sort_values(["mode", "tech", "vintage"]).reset_index(drop=True),
        pd.DataFrame(weight_records.values()), pd.DataFrame(architecture), pd.DataFrame(hydrogen), findings,
        {"files": {str(path): file_sha256(path) for path in files},
         "region": config.region, "atb_trajectory": configured_trajectory(bundle, parameter="costs"),
         "rows": len(attribution), "audit_rows_without_matched_reference": len(audit) - len(attribution),
         "legacy_currency_factor_cell": f"{config.legacy_currency_sheet}!{config.legacy_currency_cell}",
         "primary_sources": config.primary_sources,
         "identity": "current / legacy = currency factor * class-weight factor * native-basis residual",
         "applied_rpe_by_family": costs["nlr_road_retail_price_equivalent_markup"],
         "native_basis_caution": "Native amounts use different dollar years and may differ in technology, source inputs and within-class composition; this factor is not a pure engineering-price effect.",
         "counterfactual_caution": "Weight counterfactuals hold current source prices, family markup, vintage and currency fixed; they do not promote a production assumption.",
         "maximum_reconciliation_error": float(attribution.reconciliation_residual.abs().max()),
         "maximum_workbook_error": float(attribution.legacy_workbook_residual.abs().max())},
    )
