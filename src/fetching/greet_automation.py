"""Explicit copy-only Excel generation of the registered GREET vehicle-cycle bank."""

import argparse
from datetime import datetime, timezone
import json
import logging
from math import isclose, isfinite
import os
from pathlib import Path
import re
import shutil
import signal
import subprocess
import sys
import time
from uuid import uuid4

import pandas as pd
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter

from fetching.greet_vehicle_cycle import (
    COMPONENT, SOURCE, GenerationSettings, LightweightGliderAudit, ResultManifest,
    _expect, _sheet, extract_summary_rows, generation_profile, normalize_generated_results,
    registered_workbooks, result_bank_paths, validate_generation_evidence,
)
from fetching.nlr_atb_autonomie import ATB_SOURCE_ID
from utils import (
    ConfigBundle, file_sha256, load_config_bundle, load_harmonization_rules,
    resolve_artifact_path, resolve_input_path, write_dataframe_atomic,
)
from utils.files import write_text_atomic


LOGGER = logging.getLogger(__name__)


def audit_lightweight_glider(cells: dict, spec: LightweightGliderAudit) -> dict:
    """Audit the native switch against adjacent range labels; propose, never apply, a fix."""
    row = re.search(r"[0-9]+$", spec.formula_cell)[0]
    ranges = {}
    for column in spec.columns:
        label = cells.get(f"{column}{spec.header_row}")
        match = re.fullmatch(spec.header_pattern, str(label).strip())
        if match is None:
            raise ValueError(f"Unexpected GREET lightweight glider range label: {column}/{label}")
        ranges[column + row] = int(match[1])
    formula = cells.get(spec.formula_cell)
    if not isinstance(formula, str):
        raise ValueError("Missing GREET lightweight glider selection formula")
    reference = re.escape(spec.range_reference)
    pattern = r"IF\(\s*" + reference + r"\s*=\s*([0-9]+)\s*,\s*([A-Z]+[0-9]+)\s*,"
    switches = {int(miles): cell for miles, cell in re.findall(pattern, formula)}
    fallback = re.search(r",\s*([A-Z]+[0-9]+)\s*\)+\s*\+", formula)
    if len(switches) != len(ranges) - 1 or fallback is None or fallback[1] not in ranges:
        raise ValueError("Unexpected GREET lightweight glider switch layout")
    checks = [{"selected_aer_miles": miles, "referenced_glider_cell": switches.get(miles, fallback[1]),
               "referenced_glider_aer_miles": ranges[switches.get(miles, fallback[1])]}
              for miles in sorted(ranges.values())]
    # Only the condition numbers change; all glider/battery terms remain verbatim.
    proposed = re.sub(pattern, lambda m: m[0].replace(m[1], str(ranges[m[2]])), formula)
    return {"formula_cell": spec.formula_cell, "original_formula": formula,
            "range_checks": checks, "ok": all(c["selected_aer_miles"] == c["referenced_glider_aer_miles"] for c in checks),
            "proposed_formula": proposed, "correction_applied": False}


def inspect_lightweight_glider_formulas(path: Path, spec: LightweightGliderAudit) -> list[dict]:
    import openpyxl
    from openpyxl.utils.cell import column_index_from_string

    last_row = max(spec.header_row, int(re.search(r"[0-9]+$", spec.formula_cell)[0]))
    audits = []
    with path.open("rb") as handle:
        book = openpyxl.load_workbook(handle, read_only=True, data_only=False, keep_links=False)
        try:
            audits = [{"worksheet": sheet, **audit_lightweight_glider(
                _sheet(book, sheet, last_row, max(column_index_from_string(c) for c in spec.columns)), spec)}
                for sheet in spec.worksheets]
        finally:
            book.close()
    # Compare saved numeric inputs, including formula caches, without recalculation.
    row = int(re.search(r"[0-9]+$", spec.formula_cell)[0])
    with path.open("rb") as handle:
        book = openpyxl.load_workbook(handle, read_only=True, data_only=True, keep_links=False)
        try:
            for audit in audits:
                cells = _sheet(book, audit["worksheet"], last_row, max(column_index_from_string(c) for c in spec.columns))
                values = {f"{column}{row}": cells.get(f"{column}{row}") for column in spec.columns}
                if any(isinstance(value, bool) or not isinstance(value, (int, float))
                       or not isfinite(value) or value < 0 for value in values.values()):
                    raise ValueError(f"Missing or invalid GREET lightweight glider input: {audit['worksheet']}")
                audit["glider_input_values"] = values
                audit["glider_inputs_equal"] = len(set(values.values())) == 1
        finally:
            book.close()
    return audits


def probe_archetypes(bundle: ConfigBundle) -> tuple[dict[str, dict[str, int]], Path]:
    """Use current new-tech owners, efficiency archetypes and actual ATB detail labels."""
    efficiency = load_harmonization_rules(bundle, "efficiencies")
    embodied = load_harmonization_rules(bundle, "road_embodied_emissions")
    technology = pd.read_csv(resolve_input_path(bundle, "template", "technology.csv"))
    modes = [mode for mode, spec in embodied["road_classes"].items() if spec["family"] == "ldv"]
    requested = set(technology.loc[
        technology.category.isin(modes) & technology.tech.str.endswith(embodied["new_technology_suffix"]),
        "sub_category",
    ])
    atb_rules = load_harmonization_rules(bundle, "nlr_atb_autonomie")
    atb_path = resolve_input_path(bundle, "interim", atb_rules["interim_subdir"],
                                  atb_rules["components"]["vehicles"]["output_file"])
    atb = pd.read_csv(atb_path)
    if set(atb.source_id) != {ATB_SOURCE_ID}:
        raise ValueError("Unexpected registered NLR ATB archetype evidence")
    details = set(atb.loc[atb.vehicle_weight_category.eq(efficiency["atb"]["weight_categories"]["ldv"]), "vehicle_detail"])
    result = {"phev": {}, "bev": {}}
    for key in sorted(requested):
        family = embodied["powertrain_map"]["ldv"].get(key)
        if family not in {"phev", "bev"}:
            continue
        if key not in embodied["ldv_range_miles"] or key not in efficiency["ldv_archetypes"]:
            raise ValueError(f"Unsupported NLR ATB range archetype: {key}")
        miles = embodied["ldv_range_miles"][key]
        spec = efficiency["ldv_archetypes"][key]
        if family == "phev":
            if spec["range_mi"] != miles or not any(f"({miles}-mile electric range)" in d for d in details):
                raise ValueError(f"Unresolved NLR ATB PHEV range: {key}")
        elif spec["detail"] not in details or f"({miles}-mile range)" not in spec["detail"]:
            raise ValueError(f"Unresolved NLR ATB BEV range: {key}")
        result[family][key] = miles
    if not all(result.values()):
        raise ValueError("The GREET probe needs both PHEV and BEV archetypes")
    return result, atb_path


def probe_cases(rules: dict, ranges: dict[str, dict[str, int]]) -> list[dict]:
    """Cover each range, independent control changes, each truck, and reordered repeats."""
    phev, bev = sorted(set(ranges["phev"].values())), sorted(set(ranges["bev"].values()))
    ldv = sorted(rules["ldv_selectors"], key=rules["ldv_selectors"].get)
    trucks = sorted(rules["mhdv_selectors"], key=rules["mhdv_selectors"].get)

    def case(vehicle, p, b, truck):
        return {"ldv_class": vehicle, "ldv_selector": rules["ldv_selectors"][vehicle],
                "phev_range_miles": p, "bev_range_miles": b,
                "mhdv_class": truck, "mhdv_selector": rules["mhdv_selectors"][truck]}

    cases = []
    for vehicle in ldv:
        for i, distance in enumerate(bev):
            cases.append(case(vehicle, phev[i % len(phev)], distance, trucks[0]))
        for distance in phev[1:]:
            cases.append(case(vehicle, distance, bev[0], trucks[0]))
    for truck in trucks[1:]:
        cases.append(case(ldv[0], phev[0], bev[0], truck))
        cases.append(case(ldv[-1], phev[-1], bev[-1], truck))
    repeats = [cases[0], cases[len(bev) + len(phev) - 1], cases[-5], cases[-1]]
    return [dict(c, case_id=i, repeat=i >= len(cases))
            for i, c in enumerate([*cases, *reversed(repeats)])]


def _cells(sheet, last_row: int, last_column: int) -> dict:
    data = sheet.range((1, 1), (last_row, last_column)).options(ndim=2).value
    return {f"{get_column_letter(col + 1)}{row + 1}": value
            for row, values in enumerate(data) for col, value in enumerate(values)
            if value is not None}


def _cells_for(sheet, coordinates) -> dict:
    bounds = [coordinate_from_string(c) for c in coordinates]
    return _cells(sheet, max(r for _, r in bounds), max(column_index_from_string(c) for c, _ in bounds))


def _close_values(a: list[dict], b: list[dict], settings: GenerationSettings) -> bool:
    return len(a) == len(b) and all(
        isclose(x["source_value"], y["source_value"],
                rel_tol=settings.relative_tolerance, abs_tol=settings.absolute_tolerance)
        for x, y in zip(a, b, strict=True)
    )


def validate_probe_results(frame: pd.DataFrame, settings: GenerationSettings, *, cases: list[dict], rules: dict) -> dict:
    """Reject range-insensitive, order-dependent or LDV-dependent MHDV results."""
    if frame.empty or set(frame.gas_label) != set(rules["gas_rows"]) - {rules["excluded_gas"]}:
        raise ValueError("Incomplete individual-gas GREET probe results")
    keys = ["case_id", "family", "vehicle_class", "powertrain", "materials", "gas_label"]
    expected = {
        (case["case_id"], family, case[family + "_class"], block["powertrain"],
         block.get("materials", "none"), gas)
        for case in cases for family in ("ldv", "mhdv") for block in rules[family]["blocks"]
        for gas in rules["gas_rows"] if gas != rules["excluded_gas"]
    }
    actual = set(frame[keys].itertuples(index=False, name=None))
    if len(actual) != len(frame) or actual != expected:
        raise ValueError("Incomplete or duplicate GREET case/gas/material coverage")
    if set(frame.source_units) != {rules["source_units"]} or not frame.source_value.map(
        lambda x: pd.notna(x) and 0 <= x < float("inf")
    ).all():
        raise ValueError("Invalid GREET probe result values/units")
    frame = frame.copy()
    frame["range_miles"] = 0
    for powertrain in ("phev", "bev"):
        mask = frame.family.eq("ldv") & frame.powertrain.eq(powertrain)
        frame.loc[mask, "range_miles"] = frame.loc[mask, powertrain + "_range_miles"]
    keys = ["family", "vehicle_class", "powertrain", "materials", "gas_label", "range_miles"]
    comparisons = 0
    for key, group in frame.groupby(keys, dropna=False):
        reference = group.source_value.iloc[0]
        if not all(isclose(v, reference, rel_tol=settings.relative_tolerance,
                           abs_tol=settings.absolute_tolerance) for v in group.source_value):
            raise ValueError(f"Order-dependent or cross-control-dependent GREET results: {key}")
        comparisons += len(group) - 1
    response = []
    co2 = frame.loc[frame.family.eq("ldv") & frame.gas_label.eq("CO2")
                    & frame.powertrain.isin(["phev", "bev"])]
    for key, group in co2.groupby(["vehicle_class", "powertrain", "materials"]):
        by_range = group.groupby("range_miles").source_value.first()
        if len(by_range) < 2 or len(set(by_range)) != len(by_range):
            raise ValueError(f"GREET range changes did not change the vehicle-cycle result: {key}")
        response.append({"vehicle_class": key[0], "powertrain": key[1], "materials": key[2],
                         "range_values_g": {int(k): float(v) for k, v in by_range.items()}})
    return {"ok": True, "individual_gas_rows": len(frame), "independence_comparisons": comparisons,
            "range_response": response, "mhdv_classes": sorted(set(frame.loc[frame.family.eq("mhdv"), "vehicle_class"]))}


def _worker(bundle: ConfigBundle, run: Path) -> None:
    """Own a single fresh Excel process and one disposable workbook pair."""
    import xlwings as xw

    rules = load_harmonization_rules(bundle, "greet_vehicle_cycle")
    registration, paths = registered_workbooks(bundle)
    settings = registration.automation
    ranges, atb_path = probe_archetypes(bundle)
    cases = probe_cases(rules, ranges)
    hashes = {p.relative_to(bundle.repo_root).as_posix(): file_sha256(p) for p in paths}
    working = run / "working"
    working.mkdir()
    for path in paths:
        shutil.copy2(path, working / path.name)
    report = {"ok": False, "source_id": SOURCE, "component": COMPONENT, "source_hashes": hashes,
              "source_dq": bundle.sources.resolved_data_quality(SOURCE, COMPONENT).row_fields(),
              "atb_evidence": {"path": atb_path.relative_to(bundle.repo_root).as_posix(), "sha256": file_sha256(atb_path)},
              "range_archetypes": ranges, "simulation_year_changed": False, "xlwings_version": xw.__version__,
              "generation_settings": settings.model_dump(), "profile_sha256": generation_profile(bundle)}
    report["source_formula_audit"] = inspect_lightweight_glider_formulas(paths[1], settings.lightweight_glider)
    report["lightweight_bev_range_mapping_ok"] = all(f["ok"] for f in report["source_formula_audit"])
    if not report["lightweight_bev_range_mapping_ok"]:
        LOGGER.warning("Native GREET lightweight BEV glider-range formulas mismatch their labels; retained without correction; identical inputs=%s",
                       all(a["glider_inputs_equal"] for a in report["source_formula_audit"]))
    rows, exclusions, case_audit = [], [], []
    try:
        with xw.App(visible=False, add_book=False) as app:
            write_text_atomic(json.dumps({"pid": app.pid}) + "\n", run / "owned_excel.json")
            app.display_alerts = False
            app.screen_updating = False
            first = app.books.open(str(working / paths[0].name), update_links=False)
            second = app.books.open(str(working / paths[1].name), update_links=False)
            functions = []
            for check in settings.function_checks:
                book = {"greet1": first, "greet2": second}[check.workbook]
                result = book.macro(check.name)(*check.arguments)
                if result != check.expected:
                    raise ValueError(f"Unavailable or unexpected GREET VBA function: {check.name}/{result!r}")
                functions.append({**check.model_dump(), "result": result})
            report.update({"excel_version": str(app.version), "function_checks": functions,
                           "native_calculation": {"mode": app.calculation, "iteration": bool(app.api.Iteration),
                                                  "max_iterations": app.api.MaxIterations, "max_change": app.api.MaxChange}})
            before_links = list(second.api.LinkSources(1) or [])
            for link in before_links:
                if Path(link).name != paths[0].name:
                    raise ValueError(f"Unexpected GREET2 external workbook dependency: {link}")
                second.api.ChangeLink(link, str(working / paths[0].name), 1)
            after_links = list(second.api.LinkSources(1) or [])
            if not after_links or any(Path(link).resolve() != working / paths[0].name for link in after_links):
                raise ValueError("GREET2 external links are not confined to the disposable GREET1 copy")
            report["external_links"] = {"before": before_links, "after": after_links}
            if {Path(book.fullname).resolve() for book in app.books} != {working / p.name for p in paths}:
                raise ValueError("Unexpected workbook opened in the isolated GREET Excel instance")
            inputs = first.sheets[rules["inputs_sheet"]]
            for control in settings.names.values():
                native = first.names[control.name].refers_to_range
                if native.sheet.name != rules["inputs_sheet"] or native.address.replace("$", "") != control.cell:
                    raise ValueError(f"Unexpected GREET native control name: {control.name}")
            for cell, label in settings.range_labels.items():
                _expect({cell: inputs.range(cell).value}, cell, label, rules["inputs_sheet"])
            unit = settings.range_unit
            _expect({unit["cell"]: inputs.range(unit["cell"]).value}, unit["cell"], unit["label"], rules["inputs_sheet"])
            year_spec = rules["simulation_year"]
            _expect({year_spec["label_cell"]: inputs.range(year_spec["label_cell"]).value},
                    year_spec["label_cell"], year_spec["label"], rules["inputs_sheet"])
            for family, name in settings.range_dropdown_names.items():
                native_ranges = set(first.names[name].refers_to_range.value)
                if not set(ranges[family].values()) <= native_ranges:
                    raise ValueError(f"ATB ranges are unsupported by the GREET dropdown: {family}")
            ldv_inputs = second.sheets[rules["ldv"]["identity_sheet"]]
            defaults = {cell: ldv_inputs.range(cell).value for cell in settings.preserved_ldv_selectors}
            report["preserved_ldv_selectors"] = defaults
            year_cell = settings.names["simulation_year"].cell
            report["saved_simulation_year"] = inputs.range(year_cell).value
            year = registration.solved_simulation_year
            report["simulation_year_changed"] = report["saved_simulation_year"] != year
            for case in cases:
                LOGGER.info("GREET case %s: %s", case["case_id"], case)
                with app.properties(enable_events=False):
                    inputs.range(year_cell).value = year
                    for role in ("ldv_selector", "phev_range_miles", "bev_range_miles"):
                        inputs.range(settings.names[role].cell).value = case[role]
                    for family in ("ldv", "mhdv"):
                        spec = rules[family]
                        second.sheets[spec["identity_sheet"]].range(spec["selector_cell"]).value = case[family + "_selector"]
                model = {**case, "greet1_file": paths[0].relative_to(bundle.repo_root).as_posix(),
                         "greet2_file": paths[1].relative_to(bundle.repo_root).as_posix(),
                         "greet1_sha256": hashes[paths[0].relative_to(bundle.repo_root).as_posix()],
                         "greet2_sha256": hashes[paths[1].relative_to(bundle.repo_root).as_posix()],
                         "solved_simulation_year": year, "ldv_vehicle_cohort_year": registration.ldv_vehicle_cohort_year}
                previous = None
                started = time.monotonic()
                for passes in range(1, settings.max_calculation_passes + 1):
                    first.macro(settings.macro)(*settings.macro_arguments)
                    if case["case_id"] == 0 and passes == 1:
                        app.api.CalculateFullRebuild()
                    else:
                        app.api.CalculateFull()
                    current, current_excluded = [], []
                    for family in ("ldv", "mhdv"):
                        spec = rules[family]
                        identity = _cells_for(second.sheets[spec["identity_sheet"]], [spec["identity_header_cell"], spec["selector_cell"], *spec["identity_label_cells"].values()])
                        _expect(identity, spec["identity_header_cell"], spec["identity_header"], spec["identity_sheet"])
                        _expect(identity, spec["selector_cell"], case[family + "_selector"], spec["identity_sheet"])
                        selector_label = spec["identity_labels"][case[family + "_class"]]
                        _expect(identity, spec["identity_label_cells"][case[family + "_class"]], selector_label, spec["identity_sheet"])
                        layout = rules["summary_layout"]
                        last = max(b["first_row"] for b in spec["blocks"]) + len(rules["gas_rows"]) - 1
                        cells = _cells_for(second.sheets[spec["worksheet"]], [layout["scope_cell"],
                            *[f"{column}{last}" for column in (spec["column"], layout["label_column"], layout["unit_column"])]])
                        block_rows, block_exclusions = extract_summary_rows(cells, family=family, model=model, rules=rules)
                        current.extend(block_rows)
                        current_excluded.extend(block_exclusions)
                    cohort = rules["cohort_year"]
                    imports = _cells_for(second.sheets[rules["imports_sheet"]], [*settings.imported_controls.values(), cohort["label_cell"], *cohort["cells"]])
                    for key, cell in settings.imported_controls.items():
                        _expect(imports, cell, case[key], rules["imports_sheet"])
                    _expect(imports, cohort["label_cell"], cohort["label"], rules["imports_sheet"])
                    for cell in cohort["cells"]:
                        _expect(imports, cell, registration.ldv_vehicle_cohort_year, rules["imports_sheet"])
                    if inputs.range(year_cell).value != year:
                        raise ValueError("GREET simulation year changed during the fixed-year probe")
                    if any(ldv_inputs.range(c).value != v for c, v in defaults.items()):
                        raise ValueError("GREET2 default/user-defined selectors changed")
                    state = int(app.api.CalculationState)
                    if state == 1:
                        raise ValueError("Excel is still calculating after the synchronous calculation call")
                    if previous is not None and _close_values(previous, current, settings):
                        break
                    previous = current
                else:
                    raise ValueError(f"GREET vehicle-cycle totals did not converge for case {case}")
                rows.extend(current)
                exclusions.extend(current_excluded)
                case_audit.append({**case, "calculation_passes": passes, "calculation_state": state,
                                   "elapsed_seconds": time.monotonic() - started,
                                   "individual_gas_rows": len(current), "converged_vehicle_cycle_totals": True})
                write_dataframe_atomic(pd.DataFrame(rows), run / settings.files["results"])
                write_dataframe_atomic(pd.DataFrame(exclusions), run / settings.files["exclusions"])
                write_dataframe_atomic(pd.DataFrame(case_audit), run / settings.files["cases"])
                LOGGER.info("GREET case %s converged in %s passes; calculation state=%s", case["case_id"], passes, state)
            report["validation"] = validate_probe_results(pd.DataFrame(rows), settings, cases=cases, rules=rules)
            report["ok"] = True
            report["calculation_states"] = sorted({c["calculation_state"] for c in case_audit})
            report["stability_scope"] = "Individual vehicle-cycle gases and imported controls; global pending state retained as evidence"
            app.enable_events = False
            second.close()
            first.close()
    except Exception as error:
        report["error"] = f"{type(error).__name__}: {error}"
        raise
    finally:
        report["source_hashes_unchanged"] = all(file_sha256(p) == hashes[p.relative_to(bundle.repo_root).as_posix()] for p in paths)
        if not report["source_hashes_unchanged"]:
            report["ok"] = False
        write_text_atomic(json.dumps(report, sort_keys=True, indent=2) + "\n", run / settings.files["report"])


def publish_generated_results(bundle: ConfigBundle, run: Path) -> None:
    """Publish a validated complete source handoff, with its manifest written last."""
    registration, _ = registered_workbooks(bundle)
    files = registration.automation.files
    report = json.loads((run / files["report"]).read_text(encoding="utf-8"))
    frames = {key: pd.read_csv(run / files[key], float_precision="round_trip")
              for key in ("results", "exclusions", "cases")}
    validate_generation_evidence(bundle, frames["results"], frames["exclusions"], frames["cases"], report)
    bank = result_bank_paths(bundle)
    for key, name in files.items():
        write_text_atomic((run / name).read_text(encoding="utf-8"), bank[key])
    manifest = ResultManifest(schema_version=1, source_id=SOURCE, component=COMPONENT,
        profile_sha256=report["profile_sha256"], source_hashes=report["source_hashes"],
        files={bank[key].name: file_sha256(bank[key]) for key in files})
    write_text_atomic(manifest.model_dump_json(indent=2) + "\n", bank["manifest"])
    LOGGER.info("Published complete GREET result bank: %s", bank["manifest"].parent)


def generate_vehicle_cycle_results(bundle: ConfigBundle) -> Path:
    """Supervise a Windows worker; timeout cleanup targets only its fresh Excel PID."""
    if sys.platform != "win32":
        raise RuntimeError("The GREET automation probe requires Windows desktop Excel")
    registration, paths = registered_workbooks(bundle)
    settings = registration.automation
    if any(not path.is_file() for path in paths):
        raise ValueError("The GREET automation probe requires one registered existing workbook pair")
    probe_archetypes(bundle)
    run = resolve_artifact_path(bundle, "greet_automation_validation") / (
        datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "_" + uuid4().hex[:8])
    run.mkdir(parents=True)
    command = [sys.executable, "-m", "fetching.greet_automation", "--worker", "--scenario",
               str(bundle.scenario_path), "--run-dir", str(run)]
    with (run / "execution.log").open("w", encoding="utf-8") as log:
        try:
            subprocess.run(command, cwd=bundle.repo_root, stdout=log, stderr=subprocess.STDOUT,
                           check=True, timeout=settings.timeout_seconds)
        except subprocess.TimeoutExpired:
            marker = run / "owned_excel.json"
            if marker.is_file():
                os.kill(json.loads(marker.read_text())["pid"], signal.SIGTERM)
            write_text_atomic(json.dumps({"ok": False, "error": "GREET probe timeout", "timeout_seconds": settings.timeout_seconds}) + "\n",
                              run / settings.files["report"])
            raise
    report = json.loads((run / settings.files["report"]).read_text())
    if not report["ok"] or not report["source_hashes_unchanged"]:
        raise ValueError(f"GREET probe failed validation: {run}")
    publish_generated_results(bundle, run)
    normalize_generated_results(bundle)
    LOGGER.info("Validated GREET generation: %s", run)
    return run


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--scenario", default="config/scenarios/legacy_reproduction.yaml")
    parser.add_argument("--worker", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--run-dir", type=Path, help=argparse.SUPPRESS)
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    bundle = load_config_bundle(args.scenario)
    if args.worker:
        if args.run_dir is None:
            parser.error("worker requires --run-dir")
        run = args.run_dir.resolve()
        if not run.is_relative_to(resolve_artifact_path(bundle, "greet_automation_validation").resolve()):
            parser.error("worker run directory must be within its configured validation artifact route")
        _worker(bundle, run)
    else:
        print(generate_vehicle_cycle_results(bundle))


if __name__ == "__main__":
    main()
